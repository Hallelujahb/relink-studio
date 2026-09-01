import json
import os

import openpyxl
import pytest

from app.parsers import parse_file, ParseError


@pytest.fixture
def tmp_csv(tmp_path):
    path = tmp_path / "data.csv"
    path.write_text("id,name\n1,Alpha\n2,Beta\n")
    return str(path)


@pytest.fixture
def tmp_xlsx_with_trailing_blanks(tmp_path):
    path = tmp_path / "data.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["id", "name"])
    ws.append(["1", "Alpha"])
    ws.append(["2", "Beta"])
    # Simulate Excel's used-range including formatted-but-empty trailing rows.
    for r in range(4, 15):
        ws.cell(row=r, column=1).fill = openpyxl.styles.PatternFill("solid", fgColor="FFFFFF")
    wb.save(str(path))
    return str(path)


@pytest.fixture
def tmp_geojson(tmp_path):
    path = tmp_path / "data.geojson"
    data = {
        "type": "FeatureCollection",
        "features": [
            {"type": "Feature", "properties": {"id": "1", "name": "Alpha"},
             "geometry": {"type": "Point", "coordinates": [0.0, 0.0]}},
        ],
    }
    path.write_text(json.dumps(data))
    return str(path)


def test_parse_csv_returns_correct_row_count(tmp_csv):
    result = parse_file(tmp_csv, "data.csv")
    assert result["row_count"] == 2
    assert result["columns"] == ["id", "name"]
    assert result["format"] == "csv"


def test_parse_xlsx_trailing_blank_rows_bug_fixed(tmp_xlsx_with_trailing_blanks):
    """Regression test for the xlsx trailing-empty-rows bug:
    11 formatted-but-empty trailing rows must NOT inflate row_count."""
    result = parse_file(tmp_xlsx_with_trailing_blanks, "data.xlsx")
    assert result["row_count"] == 2
    assert len(result["rows"]) == 2


def test_parse_geojson_extracts_properties_and_geometry(tmp_geojson):
    result = parse_file(tmp_geojson, "data.geojson")
    assert result["row_count"] == 1
    assert result["rows"][0]["name"] == "Alpha"
    assert result["geometry"][0]["type"] == "Point"


def test_parse_unsupported_extension_raises(tmp_path):
    path = tmp_path / "data.txt"
    path.write_text("hello")
    with pytest.raises(ParseError):
        parse_file(str(path), "data.txt")


def test_parse_empty_csv_raises(tmp_path):
    path = tmp_path / "empty.csv"
    path.write_text("")
    with pytest.raises(ParseError):
        parse_file(str(path), "empty.csv")
